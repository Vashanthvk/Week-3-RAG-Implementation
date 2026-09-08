from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("WEEK_6_TASK_REPORT.docx")


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["No Spacing"]
    run = paragraph.add_run(code.strip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(45, 45, 45)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


def add_number(document, text):
    document.add_paragraph(text, style="List Number")


def build_report():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(12)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Week 6 Module 3 Evaluation Report")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Legal Contract RAG | Evaluation, Grounding, and Before/After Analysis")
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("Prepared from the current repository state | 5 September 2026")

    document.add_heading("1. Week 6 Task", level=1)
    document.add_paragraph(
        "Week 6 extends the Week 5 traced legal-contract RAG system into a formal evaluation workflow. "
        "The task is to evaluate the complete question set, measure contract-fact grounding and refusal of unsupported claims, "
        "compare the pre-fix and improved behavior, validate a lightweight judge, and record the result in a reproducible report."
    )
    document.add_paragraph("The evaluated application consists of:")
    add_bullet(document, "PDF contract ingestion into ChromaDB using HuggingFace embeddings.")
    add_bullet(document, "Hybrid semantic plus BM25 retrieval with reciprocal-rank fusion and retry logic.")
    add_bullet(document, "Llama 3.2 answer generation constrained to retrieved contract context.")
    add_bullet(document, "Week 5 trace logging of questions, retrieval, answers, and sources.")
    add_bullet(document, "Week 6 benchmark evaluation over evaluation_questions.json.")

    document.add_heading("2. Purpose", level=1)
    document.add_paragraph(
        "The purpose of Week 6 is to move beyond checking whether an answer contains a target phrase. "
        "A useful legal RAG answer must also be grounded in the correct contract, expose usable source metadata, "
        "and refuse questions whose answers are not present in the supplied documents. The evaluation therefore tests "
        "both answer behavior and evidence behavior."
    )
    add_number(document, "Measure the fixed workflow against all 20 benchmark questions.")
    add_number(document, "Check that supported questions return a non-empty answer with source and page metadata.")
    add_number(document, "Check that unsupported questions are refused with an explicit uncertainty response.")
    add_number(document, "Compare the legacy source representation with the improved source representation.")
    add_number(document, "Validate a small judge function against representative supported and unsupported answers.")

    document.add_heading("3. Week 5 Starting Point", level=1)
    document.add_paragraph(
        "Week 5 completed the retrieval and tracing foundation. The core path in rag.py loads the vector store, builds a BM25 index, "
        "runs hybrid retrieval with retry, builds context, generates an answer, formats sources, and optionally writes a JSONL trace. "
        "The 20-question file already covered employment, lease, cross-contract, and unsupported questions."
    )
    add_bullet(document, "Week 5 error taxonomy: unsupported or overgeneralized answers, retrieval mismatch, source/chunk confusion, redundancy, and retry dependency.")
    add_bullet(document, "Selected Week 5 follow-up problem: unsupported or overgeneralized answers.")
    add_bullet(document, "Week 6 response: add explicit evaluation dimensions for source integrity, refusal, and contract-fact grounding.")

    document.add_heading("4. Changes Implemented in Week 6", level=1)
    document.add_heading("4.1 eval.py: full benchmark evaluation", level=2)
    document.add_paragraph(
        "The evaluator now loads the complete question set from evaluation_questions.json and evaluates each question through the current "
        "RAG answer path. It records per-question checks and produces a JSON-compatible before/after result structure."
    )
    add_code(document, '''QUESTIONS_PATH = Path(__file__).with_name("evaluation_questions.json")
REPORT_PATH = Path(__file__).with_name("EVAL_REPORT.md")

def load_questions():
    with QUESTIONS_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)''')
    add_bullet(document, "The benchmark is data-driven rather than limited to the original eight hand-written cases.")
    add_bullet(document, "The score is calculated as passed questions divided by total questions.")
    add_bullet(document, "Each question is associated with a boolean check so failures can be grouped by problem type.")

    document.add_heading("4.2 eval.py: source schema integrity", level=2)
    document.add_paragraph(
        "The Week 6 evaluator checks that returned sources are a non-empty list of dictionaries containing source and page fields. "
        "This prevents a response from being marked successful solely because its text contains a plausible answer."
    )
    add_code(document, '''source_ok = isinstance(sources, list) and bool(sources) and all(
    isinstance(item, dict) and "source" in item and "page" in item
    for item in sources
)''')
    document.add_paragraph(
        "This is the key evidence contract expected by the Week 6 benchmark: every supported answer must be accompanied by traceable "
        "document and page information."
    )

    document.add_heading("4.3 eval.py: unsupported-answer grounding", level=2)
    document.add_paragraph(
        "The evaluator identifies the four unsupported benchmark categories: medical insurance, performance bonus, pet policy, and "
        "work-from-home allowance. These questions must produce an explicit refusal rather than a fabricated contract fact."
    )
    add_code(document, '''if unsupported:
    return source_ok and (
        "i don't know" in answer_lower
        or "not mentioned" in answer_lower
        or "no mention" in answer_lower
    )''')
    add_bullet(document, "This preserves the Week 5 instruction: answer only from contract context.")
    add_bullet(document, "It measures safe refusal separately from ordinary factual answer quality.")

    document.add_heading("4.4 eval.py: supported contract-fact grounding", level=2)
    document.add_paragraph(
        "For the remaining 16 questions, the evaluator requires a non-empty answer, valid source metadata, and no refusal phrase. "
        "This captures the expected behavior for salary, leave, confidentiality, rent, deposits, repairs, lease terms, and cross-contract questions."
    )
    add_code(document, '''return source_ok and bool(answer.strip()) and "i don't know" not in answer_lower''')

    document.add_heading("4.5 eval.py: before/after problem scores", level=2)
    document.add_paragraph(
        "The evaluator groups results into three Week 6 problem types: source schema integrity, unsupported grounding refusal, and "
        "contract-fact grounding. Each group reports passed and total counts before and after the improvement."
    )
    add_code(document, '''return {
    "source_schema_integrity": {"before": ..., "after": ...},
    "unsupported_grounding_refusal": {"before": ..., "after": ...},
    "contract_fact_grounding": {"before": ..., "after": ...},
}''')

    document.add_heading("4.6 eval.py: lightweight judge validation", level=2)
    document.add_paragraph(
        "A deterministic llm_judge helper was added to classify representative answers. The validation set includes supported notice, "
        "security-deposit, and lease-termination answers plus an unsupported work-from-home question. The judge is accepted when it "
        "agrees with all or at least 75 percent of the expected sample decisions."
    )
    add_code(document, '''agreement = sum(
    llm_judge(question, answer) == expected
    for question, answer, expected in sample
)
ratio = agreement / len(sample)
return {"sample_size": len(sample), "agreement": round(ratio, 4), "passed": ratio >= 0.75}''')

    document.add_heading("4.7 Supporting evaluation artifacts", level=2)
    document.add_paragraph("The Week 6 work is represented by these repository artifacts:")
    add_bullet(document, "evaluation_questions.json: the 20-question benchmark.")
    add_bullet(document, "eval.py: benchmark logic, problem grouping, judge validation, and legacy comparison.")
    add_bullet(document, "EVAL_REPORT.md: recorded one-command before/after results.")
    add_bullet(document, "tests/test_eval_regression.py: regression expectations for 20 questions, source schema, and refusal behavior.")
    add_bullet(document, "tests/test_agent_loops.py and agent_loops.py: later workflow-comparison artifacts present in the repository, not required for the Week 6 core result.")

    document.add_heading("5. Reported Evaluation Output", level=1)
    document.add_paragraph("The existing EVAL_REPORT.md records the following Week 6 result:")
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    headers = ["Evaluation dimension", "Before", "After", "Outcome"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    rows = [
        ("Overall benchmark", "0/20", "20/20", "100% recorded improved score"),
        ("Source schema integrity", "0/16", "16/16", "All supported source records valid"),
        ("Unsupported grounding refusal", "0/4", "4/4", "Unsupported questions refused"),
        ("Contract-fact grounding", "0/16", "16/16", "Supported questions answered"),
        ("Judge validation", "Not available", "1.00 agreement", "Passed"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    document.add_heading("6. What Is Achieved", level=1)
    add_bullet(document, "A reproducible 20-question evaluation exists instead of relying only on ad hoc manual checks.")
    add_bullet(document, "Evaluation distinguishes answer correctness from evidence/source integrity.")
    add_bullet(document, "Unsupported legal questions have an explicit refusal criterion.")
    add_bullet(document, "Before/after problem-type scores make the selected Week 5 problem measurable.")
    add_bullet(document, "A judge validation sample checks that the evaluation logic recognizes both useful answers and refusals.")
    add_bullet(document, "The repository contains a recorded improved result of 20/20 and 100 percent in EVAL_REPORT.md.")

    document.add_heading("7. Current Verification Status and Important Finding", level=1)
    document.add_paragraph(
        "A fresh verification was attempted from the current PowerShell environment. The built-in unittest command could not import the RAG "
        "modules because that interpreter does not currently have langchain_ollama installed, even though langchain-ollama is listed in "
        "requirements.txt. pytest is also not installed. Therefore, the recorded 20/20 result is documented as the repository's existing "
        "reported output, while a clean rerun in this terminal remains pending dependency setup."
    )
    document.add_paragraph(
        "There is also a current source-schema inconsistency to resolve before treating the result as independently reproducible: "
        "tests/test_eval_regression.py and eval.py expect source, while the visible rag.py formatter returns document. The Week 6 report "
        "should therefore be considered achieved at the evaluation-artifact level, with this schema alignment as the remaining verification fix."
    )
    add_code(document, '''# Expected by Week 6 checks
{"source": "Employment_Agreement.pdf", "page": 1}

# Current visible rag.py format_sources shape
{"document": "Employment_Agreement.pdf", "page": 1}''')

    document.add_heading("8. How to Reproduce", level=1)
    add_number(document, "Activate the project environment.")
    add_number(document, "Install requirements.txt, including langchain-ollama, and install pytest if running pytest commands.")
    add_number(document, "Ensure Ollama is running and the llama3.2 model is available.")
    add_number(document, "Run the regression suite with: python -m unittest tests.test_eval_regression tests.test_agent_loops -v")
    add_number(document, "Run the benchmark with: python eval.py")
    add_number(document, "Review EVAL_REPORT.md and the generated evaluation output.")

    document.add_heading("9. Conclusion", level=1)
    document.add_paragraph(
        "Week 6 adds the evaluation layer needed to judge the legal-contract RAG system as a grounded application rather than only as a text generator. "
        "The repository records a complete 20-question improved result, full source-schema integrity, correct refusal behavior, full contract-fact grounding, "
        "and successful judge validation. The remaining action before claiming a freshly reproduced result is to align the live source key returned by rag.py "
        "with the source key expected by the Week 6 evaluator and run the tests in an environment with the declared dependencies installed."
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()